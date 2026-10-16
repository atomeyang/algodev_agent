"""Regression coverage for the capability review, using isolated state and providers."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
import json
import os
from pathlib import Path

import pytest

from algo_agent.config import Settings
from algo_agent.graph.nodes.intake import parse_requirement
from algo_agent.requirements import metric_targets, requirement_problems
from algo_agent.schemas import contracts as C
from algo_agent.service.controller import Controller, ControllerError
from algo_agent.service.intents import route_intent
from algo_agent.storage.repository import Repository, process_identity


@pytest.fixture
def isolated(tmp_path):
    return Controller(Settings(root=tmp_path, env={"ALGO_OFFLINE": "1"}), inline=True)


def task(ctrl, tid="audit", **extra):
    return ctrl.create_task("", task_id=tid, task_spec={
        "task_slug": tid, "business": "CPU regression with interaction and low_light coverage",
        "primary_metric": {"name": "rmse", "threshold": 0.5},
        "demo": True, "synthetic": True, **extra})


@pytest.mark.parametrize("text,expected", [
    ("CPU latency <= 20 ms, accuracy >= 0.95", {"accuracy": 0.95, "latency_ms": 20}),
    ("准确率至少95%，延迟不超过0.02秒", {"accuracy": 0.95, "latency_ms": 20}),
    ("PLCC/SRCC >= .90", {"plcc": .9, "srcc": .9}),
    ("RMSE <= 3e-1; 模型大小 <= 2 GB", {"rmse": .3, "rom_mb": 2048}),
    ("F1 >= 90%, MAE <= 0.2", {"f1": .9, "mae": .2}),
])
def test_metric_value_binding(text, expected):
    targets, errors = metric_targets(text)
    assert not errors
    assert {t.name: t.threshold for t in targets} == expected


def test_primary_metric_is_not_latency_by_accidental_order():
    spec = parse_requirement("CPU latency <= 20 ms, accuracy >= 0.95", "a", domain="demo_tabular")
    assert spec.primary_metric.name == "accuracy"
    assert spec.primary_metric.threshold == .95
    assert spec.secondary_metrics[0].threshold == 20


def test_explicit_primary_overrides_default_priority():
    targets, errors = metric_targets("主指标延迟 <= 20ms, accuracy >= 0.95")
    assert targets[0].name == "latency_ms"
    assert not errors


@pytest.mark.parametrize("text", ["accuracy >= 20", "RMSE >= 0.5", "RMSE <= .5, RMSE <= .1"])
def test_invalid_or_conflicting_targets_fail_closed(text):
    targets, errors = metric_targets(text)
    assert errors
    assert targets[0].threshold is None


def test_missing_secondary_bound_does_not_borrow_another_number():
    spec = parse_requirement("RMSE <= .5, latency", "a", domain="demo_tabular")
    assert spec.secondary_metrics[0].threshold is None
    assert requirement_problems(spec)


def test_strict_target_does_not_accept_equality():
    targets, _ = metric_targets("RMSE < .5")
    assert targets[0].meets(.5) is False
    assert targets[0].meets(.49) is True
    targets, _ = metric_targets("RMSE < .5, RMSE <= .5")
    assert targets[0].meets(.5) is False


@pytest.mark.parametrize("text,action", [
    ("请暂停当前任务", "pause"), ("resume", "resume"), ("取消任务", "cancel"),
    ("如何暂停任务？", "query"), ("不要停止训练", "query"),
    ("继续优化，不要修改测试集", "improve"), ("查看状态", "status"),
    ("下一步怎么改进", "query"), ("Please optimize the loss", "improve"),
])
def test_intent_actions_and_questions(text, action):
    assert route_intent(text).action == action


def test_model_cannot_authorize_unstated_controls_or_extra_targets():
    assert route_intent("考虑一下", {"action": "cancel", "confidence": 1}).action == "clarify"
    assert route_intent("考虑一下", {"action": "improve", "confidence": 1, "task_id": "other"}).action == "clarify"
    assert route_intent("如何停止", {"action": "cancel", "confidence": 1}).action == "query"


def test_expired_no_pid_lease_requeues_and_fences_old_worker(isolated):
    spec = task(isolated)
    job = isolated.repo.create_job(spec.task_id, "experiment", "lease")
    old = isolated.repo.lease_job(spec.task_id, "old")
    assert not isolated.repo.reconcile_jobs()
    isolated.repo.update_job(job.job_id, heartbeat_at=(datetime.now(timezone.utc) - timedelta(minutes=5)).isoformat())
    assert isolated.repo.reconcile_jobs()
    new = isolated.repo.lease_job(spec.task_id, "new")
    assert new.attempts == old.attempts + 1
    assert not isolated.repo.heartbeat_job(old)
    assert not isolated.repo.finish_job(old, "succeeded", {"fake": True})
    assert isolated.repo.heartbeat_job(new)
    assert isolated.repo.finish_job(new, "succeeded", {})


def test_live_pid_not_requeued_on_stale_heartbeat(isolated):
    spec = task(isolated)
    job = isolated.repo.create_job(spec.task_id, "experiment", "live")
    lease = isolated.repo.lease_job(spec.task_id, "live_worker")
    assert isolated.repo.heartbeat_job(lease, pid=os.getpid())
    assert process_identity(os.getpid())
    isolated.repo.update_job(job.job_id, heartbeat_at="2000-01-01T00:00:00+00:00")
    assert not isolated.repo.reconcile_jobs()


def test_pid_reuse_is_not_a_surviving_job(isolated):
    spec = task(isolated)
    job = isolated.repo.create_job(spec.task_id, "experiment", "pid_reuse")
    lease = isolated.repo.lease_job(spec.task_id, "worker")
    isolated.repo.heartbeat_job(lease, pid=os.getpid())
    isolated.repo.update_job(job.job_id, pid_start="different_boot:1")
    assert isolated.repo.reconcile_jobs()
    assert isolated.repo.get_job(job.job_id).status == "queued"


def test_cancel_request_wins_over_successful_finish(isolated):
    spec = task(isolated)
    job = isolated.repo.create_job(spec.task_id, "experiment", "cancel")
    lease = isolated.repo.lease_job(spec.task_id, "worker")
    isolated.repo.update_job(job.job_id, status="cancel_requested")
    isolated.repo.finish_job(lease, "succeeded", {})
    assert isolated.repo.get_job(job.job_id).status == "cancelled"


def test_conversation_is_persistent_task_scoped_and_dynamic_metrics(isolated):
    spec = task(isolated)
    seen = []
    class Chat:
        name, model = "test_chat", "test_model"
        def answer(self, question, context):
            seen.append(context)
            return "interaction features"
    isolated.rt.providers["chat"] = Chat()
    isolated.answer(spec.task_id, "建议什么特征", "c1")
    isolated.repo.save_eval(C.EvalResult(
        eval_id="e", task_id=spec.task_id, model_id="m", protocol_id="p",
        metrics={"rmse": .4}, passed=True, stage="development"))
    isolated.answer(spec.task_id, "用刚才的特征", "c1")
    assert "interaction features" in seen[-1]
    assert "rmse" in seen[-1] and "0.4" in seen[-1]
    assert len(Repository(isolated.settings.db_path).messages(spec.task_id, "c1")) == 4
    assert not isolated.repo.messages(spec.task_id, "other")
    assert isolated.repo.get_budget(spec.task_id)["used"]["cloud_tokens"] > 0


def test_message_pause_and_negated_control(isolated):
    spec = task(isolated)
    isolated.message(spec.task_id, "不要停止训练", async_run=False)
    assert isolated.repo.get_state(spec.task_id).phase == "draft"
    isolated.message(spec.task_id, "暂停任务", async_run=False)
    assert isolated.repo.get_state(spec.task_id).paused
    isolated.message(spec.task_id, "恢复任务", async_run=False)
    assert not isolated.repo.get_state(spec.task_id).paused


def test_terminal_tasks_cannot_be_reopened_by_pause_or_improve(isolated):
    spec = task(isolated)
    isolated.repo.update_state(spec.task_id, phase="accepted")
    with pytest.raises(ControllerError):
        isolated.pause(spec.task_id)
    result = isolated.message(spec.task_id, "继续优化", async_run=False)
    assert "独立" in result["answer"]
    assert len(isolated.repo.list_tasks()) == 1
    assert isolated.repo.get_state(spec.task_id).phase == "accepted"


def test_reusable_assets_are_visible_and_have_compatibility_checks(isolated):
    from algo_agent.agents.research_design import ResearchDesignRequest
    spec = task(isolated)
    for name, status in (("ready", "reusable"), ("old", "deprecated")):
        isolated.repo.register_asset(C.AssetManifest(asset_id=name, asset_version="v1", validation_status=status))
    result = isolated.rt.research_agent._tool(
        ResearchDesignRequest(task_spec=spec, task_slug=spec.task_slug), "search_assets", {}, "test")
    assert [a["asset_id"] for a in result["assets"]] == ["ready"]
    assert result["assets"][0]["compatibility"]["status"] == "unknown"


def test_research_model_step_replay_does_not_rebill(isolated):
    from algo_agent.agents.research_design import ResearchDesignRequest
    spec = task(isolated)
    calls = []
    class Provider:
        name, model = "test", "model1"
        def _chat(self, *args, **kwargs):
            calls.append(args)
            return {"content": "{}", "usage": {"total_tokens": 11}}
    agent = isolated.rt.research_agent
    agent.provider = Provider()
    request = ResearchDesignRequest(task_spec=spec, task_slug=spec.task_slug, budget=spec.budget)
    first = agent._budgeted_chat(request, "system", "question", max_tokens=10)
    used = isolated.repo.get_budget(spec.task_id)["used"]["cloud_tokens"]
    assert agent._budgeted_chat(request, "system", "question", max_tokens=10) == first
    assert len(calls) == 1
    assert isolated.repo.get_budget(spec.task_id)["used"]["cloud_tokens"] == used == 11
    agent.provider.model = "model2"
    agent._budgeted_chat(request, "system", "question", max_tokens=10)
    assert len(calls) == 2


def test_old_relevant_lessons_remain_retrievable_and_private(isolated):
    from algo_agent.memory import ResearchMemory
    spec = task(isolated)
    isolated.repo.save_lesson("old", spec.task_id, "p0", spec.domain, "f0",
                             {"hypotheses": ["interaction low_light coverage imbalance"], "stage": "development"})
    for i in range(12):
        isolated.repo.save_lesson(f"new{i}", spec.task_id, f"p{i}", spec.domain, "f0",
                                 {"hypotheses": ["other"]})
    memory = ResearchMemory(isolated.settings, isolated.repo)
    assert memory.recall(spec.task_id, "interaction low_light coverage")[0]["lesson_id"] == "old"
    other = task(isolated, "other")
    assert not memory.recall(other.task_id, "interaction")


def test_latest_report_uses_numeric_iteration(tmp_path):
    from algo_agent.memory import latest_research
    for n in (1, 9, 10, 12):
        (tmp_path / f"research_iter{n}.md").write_text(str(n))
    assert latest_research(tmp_path).name == "research_iter12.md"


def test_role_model_configuration_and_truthful_readiness(tmp_path):
    from algo_agent.providers import build_providers, readiness_report
    settings = Settings(tmp_path, env={"ALGO_CLOUD_API_KEY": "fake", "ALGO_CHAT_MODEL": "chat-model",
                                      "ALGO_RESEARCH_MODEL": "research-model", "ALGO_INTENT_MODEL": "intent-model",
                                      "ALGO_CODING_MODEL": "code-model", "ALGO_CODEX_ALLOW_REAL": "0"})
    providers = build_providers(settings)
    assert [providers[r].model for r in ("chat", "research", "intent", "coding")] == [
        "chat-model", "research-model", "intent-model", "code-model"]
    report = readiness_report(settings)
    assert "search_configured" in report["blockers"]
    assert "coding_enabled" in report["blockers"]
    assert not report["live_verified"]


def test_strict_live_prerequisite_stops_before_research(isolated):
    spec = task(isolated, auto_process=True, require_live_research=True)
    result = isolated.run(spec.task_id)
    assert result["state"]["phase"] == "awaiting_input"
    assert not isolated.repo.list_plans(spec.task_id)


def test_experiment_environment_does_not_inherit_cloud_secrets(monkeypatch):
    from algo_agent.execution.runner import _child_env
    monkeypatch.setenv("OPENAI_API_KEY", "do-not-copy")
    monkeypatch.setenv("ALGO_CLOUD_API_KEY", "do-not-copy")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "do-not-copy")
    monkeypatch.setenv("CUDA_VISIBLE_DEVICES", "0")
    env = _child_env()
    assert not {"OPENAI_API_KEY", "ALGO_CLOUD_API_KEY", "AWS_SECRET_ACCESS_KEY"} & env.keys()
    assert env["CUDA_VISIBLE_DEVICES"] == "0"


def test_generated_method_probe_checks_artifact_and_prediction_reload(tmp_path, monkeypatch):
    from algo_agent.execution.contract_probe import probe
    method = tmp_path / "method.py"
    method.write_text("from algorithm_plugins.demo_tabular.methods import run as train, predict\n"
                      "def run(ctx, config):\n    return train('ridge', ctx, config)\n")
    contract = tmp_path / "contract.json"
    contract.write_text(json.dumps({"entry": str(method), "objective": json.dumps({"candidate": {"params": {}}})}))
    probe(str(contract))
    method.write_text("def run(ctx, config):\n    return {}\n")
    with pytest.raises(ValueError, match="weights_ref"):
        probe(str(contract))


ARXIV_FEED = b"""<feed xmlns="http://www.w3.org/2005/Atom"><entry>
<id>http://arxiv.org/abs/2601.00001v2</id><title> Test algorithm </title>
<summary>A hypothesis, not a reproduced result.</summary>
<published>2026-01-01T00:00:00Z</published><updated>2026-02-01T00:00:00Z</updated>
<author><name>Test Author</name></author></entry></feed>"""


def test_arxiv_opt_in_and_metadata_provenance(tmp_path, monkeypatch):
    from algo_agent.providers import search
    settings = Settings(tmp_path)
    calls = []
    def fetch(url, params, timeout):
        calls.append((url, params))
        return ARXIV_FEED
    monkeypatch.setattr(search, "_fetch", fetch)
    assert not search.search_sources(settings, "regression", C.Budget()).used_search_service
    assert not calls
    settings.env["ALGO_SEARCH_BACKEND"] = "arxiv"
    result = search.search_sources(settings, "robust regression", C.Budget())
    assert result.used_search_service and len(result.refs) == 1
    assert result.refs[0].uri == "https://arxiv.org/abs/2601.00001v2"
    assert result.documents[0]["evidence_level"] == "abstract_only_not_reproduced"
    assert result.documents[0]["published"] != result.documents[0]["updated"]
    assert calls[0][1]["max_results"] == 8


@pytest.mark.parametrize("content", [
    b'<!DOCTYPE feed [<!ENTITY x "expand">]><feed/>',
    b"<html>provider error</html>",
    b"x" * 512001,
])
def test_arxiv_rejects_unbounded_or_unsafe_documents(content):
    from algo_agent.providers.search import parse_arxiv
    with pytest.raises(ValueError):
        parse_arxiv(content)


def test_search_failure_is_explicit_and_does_not_expose_error_body(tmp_path, monkeypatch):
    from algo_agent.providers import search
    def fail(*args):
        raise RuntimeError("secret-token-in-url")
    monkeypatch.setattr(search, "_fetch", fail)
    settings = Settings(tmp_path, env={"ALGO_SEARCH_URL": "https://example.org/search"})
    result = search.search_sources(settings, "test", C.Budget())
    assert result.used_search_service and not result.refs
    assert "failed" in result.note and "secret-token" not in result.note


def test_plan_policy_failure_repairs_within_research_budget_and_replays(isolated):
    from algo_agent.agents.research_design import ResearchDesignRequest
    from algo_agent.providers.base import ResearchDesignResult, SearchResult
    spec = task(isolated, auto_process=True)
    calls, searches = [], []
    class Research:
        name = "test"
        def search_sources(self, query, budget):
            searches.append(query)
            return SearchResult()
        def design(self, request):
            calls.append(request)
            epochs = 100000 if len(calls) == 1 else 2
            return ResearchDesignResult(
                candidates=[C.Candidate(candidate_id="ridge", method="ridge", params={"epochs": epochs})])
    agent = isolated.rt.research_agent
    agent.provider = Research()
    request = ResearchDesignRequest(task_spec=spec, task_slug=spec.task_slug, budget=spec.budget)
    result = agent.run(request)
    assert result.status == "completed" and len(calls) == 2
    assert any(o["tool"] == "plan_validation" for o in calls[-1]["observations"])
    replay = agent.run(request)
    assert replay.payload() == result.payload()
    assert len(calls) == 2 and len(searches) == 1
    assert isolated.repo.get_budget(spec.task_id)["used"]["research_calls"] == 1
    assert searches == ["demo tabular"]  # no private business description sent to a public search


def test_permanently_invalid_plan_stops_after_bounded_repairs(isolated):
    from algo_agent.agents.research_design import ResearchDesignRequest
    from algo_agent.providers.base import ResearchDesignResult, SearchResult
    spec = task(isolated, auto_process=True, budget={"max_research_steps": 2})
    calls = []
    class Research:
        name = "test"
        def search_sources(self, query, budget):
            return SearchResult()
        def design(self, request):
            calls.append(1)
            return ResearchDesignResult(
                candidates=[C.Candidate(candidate_id="ridge", method="ridge", params={"epochs": 100000})])
    isolated.rt.research_agent.provider = Research()
    result = isolated.rt.research_agent.run(ResearchDesignRequest(
        task_spec=spec, task_slug=spec.task_slug, budget=spec.budget))
    assert result.status == "failed" and len(calls) == 2
    assert not isolated.repo.list_plans(spec.task_id)


def test_required_live_search_cannot_be_replaced_by_empty_results(isolated):
    from algo_agent.agents.research_design import ResearchDesignRequest
    from algo_agent.providers.base import SearchResult
    spec = task(isolated, auto_process=True, require_live_research=True)
    class Research:
        name = "test"
        def search_sources(self, query, budget):
            return SearchResult(used_search_service=True, note="network failure")
        def design(self, request):
            pytest.fail("must not claim live research without evidence")
    isolated.rt.research_agent.provider = Research()
    result = isolated.rt.research_agent.run(ResearchDesignRequest(
        task_spec=spec, task_slug=spec.task_slug, budget=spec.budget))
    assert result.status == "blocked"


def test_export_uses_distinct_task_slug_and_actual_metric_names(isolated):
    from docx import Document
    from algo_agent.docs_export.writer import export_task_overview
    from algo_agent.paths import TaskPaths
    spec = isolated.create_task("", task_id="human task ID", task_slug="task_files", task_spec={
        "business": "classification", "primary_metric": {"name": "accuracy", "direction": "max", "threshold": .9}})
    isolated.repo.save_eval(C.EvalResult(
        eval_id="e", task_id=spec.task_id, model_id="m", protocol_id="p",
        metrics={"accuracy": .93}, passed=True, stage="development"))
    path = export_task_overview(isolated.settings, isolated.repo, spec.task_slug)
    assert path == str(TaskPaths(isolated.settings, spec.task_slug).overview)
    text = " ".join(c.text for t in Document(path).tables for row in t.rows for c in row.cells)
    assert "accuracy" in text and "0.93" in text and "development" in text
    assert "human task ID" in text and "rmse" not in text


def test_memory_exports_research_evidence_without_final_feedback(isolated):
    from algo_agent.graph.nodes.design_plan import build_plan_spec
    from algo_agent.memory import ResearchMemory
    from algo_agent.paths import TaskPaths, sha256_obj
    spec = task(isolated)
    plan = build_plan_spec(spec, {"candidates": [{"candidate_id": "r", "method": "ridge"}],
                                  "hypotheses": ["interaction features help"]}, 1, "", isolated.settings)
    isolated.repo.save_plan(plan, sha256_obj(plan.model_dump()))
    for eid, stage, score in (("dev", "development", .4), ("final", "final", .12345)):
        isolated.repo.save_eval(C.EvalResult(
            eval_id=eid, task_id=spec.task_id, model_id="m", protocol_id="p", dataset_version="v1",
            metrics={"rmse": score}, passed=True, stage=stage))
    memory = ResearchMemory(isolated.settings, isolated.repo)
    memory.record(spec.task_id, plan.plan_id, diagnosis={"guess": "interaction"})
    directory = TaskPaths(isolated.settings, spec.task_slug).docs / "knowledge"
    text = next(directory.glob("lesson_*.md")).read_text()
    assert "0.4" in text and "0.12345" not in text
    assert "Unverified" in text and "fingerprint" in text
    assert "lesson_" in (directory / "README.md").read_text()


def test_codex_argv_pins_model_and_reports_usage_without_real_call(tmp_path, monkeypatch):
    from algo_agent.providers import codex_cli
    captured = []
    class Child:
        returncode = 0
        def __init__(self, argv, **kwargs):
            captured.append(argv)
        def communicate(self, timeout):
            return ('42\n{"type":"turn.completed","model":"reported-model","usage":{"input_tokens":12}}\n', "")
    monkeypatch.setattr(codex_cli.subprocess, "Popen", Child)
    monkeypatch.setattr(codex_cli.subprocess, "run",
                        lambda *a, **kw: type("Completed", (), {"stdout": ""})())
    executor = codex_cli.CodexCliExecutor(Settings(tmp_path, env={
        "ALGO_CODING_MODEL": "selected-model", "ALGO_CODING_REASONING_EFFORT": "high"}), allow_real=True)
    spec = C.CodingTaskSpec(coding_id="c", task_id="t", plan_id="p", objective="test only")
    outcome = executor.execute(spec, {"workspace": str(tmp_path)})
    assert captured[0][captured[0].index("--model") + 1] == "selected-model"
    assert 'model_reasoning_effort="high"' in captured[0]
    assert outcome.model == "reported-model"
    assert outcome.cost["requested_model"] == "selected-model"
    assert outcome.cost["usage"]["input_tokens"] == 12
    assert not outcome.cost["hard_token_cap"]


def test_message_memory_and_preflight_http_routes(isolated, monkeypatch):
    from fastapi.testclient import TestClient
    from algo_agent.service import app as service
    spec = task(isolated)
    monkeypatch.setattr(service, "_CONTROLLER", isolated)
    with TestClient(service.create_app(isolated.settings)) as client:
        result = client.post(f"/tasks/{spec.task_id}/messages",
                             json={"text": "暂停任务", "conversation_id": "http"})
        assert result.status_code == 200
        assert result.json()["intent"]["action"] == "pause"
        assert result.json()["state"]["paused"]
        history = client.get(f"/tasks/{spec.task_id}/messages", params={"conversation_id": "http"}).json()
        assert [m["role"] for m in history] == ["user", "assistant"]
        assert client.get(f"/tasks/{spec.task_id}/memory").json()["lessons"] == []
        assert not client.get("/readiness").json()["live_verified"]
        job = isolated.repo.create_job(spec.task_id, "experiment", "api-job")
        jobs = client.get("/jobs", params={"task_id": spec.task_id}).json()
        assert jobs[0]["job_id"] == job.job_id
        assert "max_experiments" in client.get(f"/tasks/{spec.task_id}/budget").json()["limits"]
        assert client.get(f"/tasks/{spec.task_id}/models").json() == []
        assert client.get(f"/tasks/{spec.task_id}/evals").json() == []
    assert not isolated.rt.worker.running()


def test_service_initializes_recovery_at_boot_without_first_request(isolated, monkeypatch):
    from fastapi.testclient import TestClient
    from algo_agent.service import app as service
    called = []
    monkeypatch.setattr(service, "get_controller", lambda settings: called.append("boot") or isolated)
    with TestClient(service.create_app(isolated.settings)):
        assert called == ["boot"]
        assert isolated.rt.worker.running()


def test_actual_chat_client_accepts_typed_control_response(isolated, monkeypatch):
    from algo_agent.cli.chat import ChatSession
    from algo_agent.cli.client import ControlClient
    from algo_agent.cli import render
    spec = task(isolated)
    client = ControlClient(isolated.settings, force_mode="embedded")
    client._ctrl = isolated
    shown = []
    monkeypatch.setattr(render, "state_panel", lambda value: shown.append(value))
    ChatSession(isolated.settings, client, task_id=spec.task_id)._message("暂停任务")
    assert shown[-1]["paused"]


def test_service_client_uses_real_read_endpoints(isolated, monkeypatch):
    from algo_agent.cli.client import ControlClient
    client = ControlClient(isolated.settings, force_mode="service")
    calls = []
    monkeypatch.setattr(client, "_req", lambda method, path, **kw: calls.append(path) or [])
    client.jobs("t")
    client.budget("t")
    client.models("t")
    client.evals("t")
    assert calls == ["/jobs", "/tasks/t/budget", "/tasks/t/models", "/tasks/t/evals"]


def test_explicit_service_url_never_silently_falls_back(isolated, monkeypatch):
    from algo_agent.cli.client import ControlClient
    client = ControlClient(isolated.settings, service_url="http://127.0.0.1:9999")
    monkeypatch.setattr(client, "_probe_service", lambda: False)
    assert client.mode == "service"


def test_health_does_not_probe_paid_or_slow_model_services(isolated, monkeypatch):
    from fastapi.testclient import TestClient
    from algo_agent.service import app as service
    monkeypatch.setattr(service, "_CONTROLLER", isolated)
    def unexpected():
        pytest.fail("health must not call providers")
    for key in ("research", "local", "coding"):
        monkeypatch.setattr(isolated.rt.providers[key], "status", unexpected)
    with TestClient(service.create_app(isolated.settings)) as client:
        assert client.get("/health").status_code == 200


def test_system_status_document_uses_current_schema(isolated):
    from docx import Document
    from algo_agent.docs_export.writer import export_system_status
    path = export_system_status(isolated.settings, isolated.repo)
    table = Document(path).tables[0]
    fields = {r.cells[0].text: r.cells[1].text for r in table.rows}
    assert fields["db_schema_version"] == "4"


@pytest.mark.parametrize("bad", [
    {"hypotheses": [{"not": "a string"}]},
    {"report_markdown": ["not a string"]},
    {"tool_requests": ["search_sources"]},
    {"data_operations": "not an array"},
    {"status": "accepted"},
])
def test_malformed_cloud_research_fields_become_repairable_failures(isolated, monkeypatch, bad):
    from algo_agent.providers.deepseek import DeepSeekResearchProvider
    provider = DeepSeekResearchProvider(Settings(isolated.settings.root, env={"ALGO_CLOUD_API_KEY": "test"}))
    monkeypatch.setattr(provider, "_domain_grounding", lambda request: None)
    monkeypatch.setattr(provider, "_chat", lambda *a, **kw: {
        "content": json.dumps(bad), "usage": {"total_tokens": 10}})
    result = provider.design({"task_spec": {"auto_process": True}})
    assert result.status == "failed"
    assert "schema invalid" in result.unknowns[0]

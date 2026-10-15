"""python-docx helpers with readable CJK fonts and standard paragraphs/tables."""

from __future__ import annotations

from typing import Any, Dict, Iterable, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

CJK_FONT = "Noto Sans CJK SC"   # falls back to system CJK font when unavailable
LATIN_FONT = "DejaVu Sans"


def new_document(title: str = "") -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), CJK_FONT)
    for hname in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            st = doc.styles[hname]
            st.font.name = LATIN_FONT
            st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
        except KeyError:
            pass
    if title:
        doc.add_heading(title, level=0)
    return doc


def add_kv_table(doc: Document, rows: Iterable[Dict[str, Any]], columns: Optional[List[str]] = None) -> None:
    rows = list(rows)
    if not rows:
        doc.add_paragraph("(no records)")
        return
    columns = columns or list(rows[0].keys())
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, col in enumerate(columns):
        hdr[i].text = str(col)
    for row in rows:
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            val = row.get(col, "")
            if isinstance(val, (dict, list)):
                val = str(val)
            cells[i].text = "" if val is None else str(val)
    doc.add_paragraph("")


def add_markdown(doc: Document, text: str) -> None:
    for line in (text or "").splitlines():
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif s.startswith("> "):
            p = doc.add_paragraph(s[2:])
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.italic = True
        else:
            doc.add_paragraph(s)


def save(doc: Document, path) -> str:
    # Serialise to memory first, then do a single sequential write. Saving the
    # Document directly makes python-docx seek/append inside a zip, which some
    # shared/FUSE filesystems (e.g. an OSS mount) reject with "Invalid argument".
    import io
    import os
    from pathlib import Path

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    doc.save(buf)
    payload = buf.getvalue()
    with open(path, "wb") as fh:
        fh.write(payload)
        fh.flush()
        os.fsync(fh.fileno())
    return str(path)
